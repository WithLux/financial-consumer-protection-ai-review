import io
import unittest
from docx import Document
from fastapi.testclient import TestClient
import test_backend as helpers

FIELDS=dict(applicant_department='虚构零售产品部',reviewer='测试审查员',review_date='2026-10-30',review_scope='虚构理财营销材料的文本审查',final_conclusion='经人工复核，需修改确认问题后再送审。',remediation_notes='请业务部门逐项核实并修改，本意见不替代正式审批。')

def all_text(document):
    return '\n'.join([p.text for p in document.paragraphs]+[c.text for t in document.tables for row in t.rows for c in row.cells])

class ClosureTests(unittest.TestCase):
    setUp=helpers.BackendTests.setUp
    tearDown=helpers.BackendTests.tearDown
    upload=helpers.BackendTests.upload
    review=helpers.BackendTests.review

    def test_human_changes_report_snapshot_and_isolation(self):
        r=self.review(self.upload('产品保证收益。虚构独有待驳回语句。风险和费用已说明。')['document_id'])
        rid=r['review_id'];issue=r['issues'][0];initial=issue.copy()
        edited=self.client.patch(f'/api/reviews/{rid}/issues/{issue["issue_id"]}',json=dict(expected_version=r['version'],review_status='edited',decision_reason='结合背景修订',severity='low',reason='人工修订说明',suggestion='人工修订建议'),headers=self.headers)
        self.assertEqual(edited.status_code,200);version=edited.json()['version']
        added=self.client.post(f'/api/reviews/{rid}/issues',json={'expected_version':version,'decision_reason':'人工核实新增','issue':dict(risk_type='other',severity='medium',original_text='保证收益',location=issue['location'],reason='人工新增独有说明',regulation_ids=['REG-NET-10-3'],suggestion='人工新增建议')},headers=self.headers)
        self.assertEqual(added.status_code,201);version=added.json()['version']
        rejected=self.client.post(f'/api/reviews/{rid}/issues',json={'expected_version':version,'decision_reason':'先补充核实','issue':dict(risk_type='other',severity='low',original_text='虚构独有待驳回语句',location={'segment_id':'s1','start':7,'end':16},reason='不应进入正式文书的说明',regulation_ids=[],suggestion='不应进入正式文书的建议')},headers=self.headers)
        version=rejected.json()['version'];rejected_id=rejected.json()['issue']['issue_id']
        patch=self.client.patch(f'/api/reviews/{rid}/issues/{rejected_id}',json=dict(expected_version=version,review_status='rejected',decision_reason='不成立，驳回'),headers=self.headers);version=patch.json()['version']
        self.assertEqual(self.client.post(f'/api/reviews/{rid}/finalize',json=dict(expected_version=version,coverage_confirmed=True,limitations_acknowledged=True),headers=self.headers).status_code,409)
        saved=self.client.patch(f'/api/reviews/{rid}/report-fields',json={'expected_version':version,'fields':FIELDS},headers=self.headers)
        self.assertEqual(saved.status_code,200,saved.text);version=saved.json()['version']
        current=self.client.get(f'/api/reviews/{rid}').json()
        self.assertEqual(current['original_issues'][0]['reason'],initial['reason'])
        self.assertEqual(current['issues'][0]['severity'],'low')
        self.assertEqual(current['summary']['manual_added'],2)
        self.assertEqual(current['summary']['formal_total'],2)
        final=self.client.post(f'/api/reviews/{rid}/finalize',json=dict(expected_version=version,coverage_confirmed=True,limitations_acknowledged=True),headers=self.headers)
        self.assertEqual(final.status_code,200,final.text)
        report=self.client.get(f'/api/reviews/{rid}/report')
        doc=Document(io.BytesIO(report.content));text=all_text(doc)
        for value in ['金融消费者权益保护审查意见书',FIELDS['reviewer'],FIELDS['applicant_department'],FIELDS['final_conclusion'],'人工修订说明','人工新增独有说明','REG-NET-10-3']:
            self.assertIn(value,text)
        self.assertNotIn('虚构独有待驳回语句',text)
        self.assertNotIn('不应进入正式文书',text)
        self.assertIn(current['regulations']['REG-NET-10-3']['source_url'],text)
        locked=self.client.patch(f'/api/reviews/{rid}/report-fields',json={'expected_version':final.json()['version'],'fields':FIELDS},headers=self.headers)
        self.assertEqual(locked.status_code,409)
        other=TestClient(self.app,base_url='https://testserver');s=other.post('/api/session',json={'access_token':'test-secret'}).json()
        self.assertEqual(other.get(f'/api/reviews/{rid}').status_code,404)
        self.assertEqual(other.get(f'/api/reviews/{rid}/report').status_code,404)
        self.assertEqual(other.patch(f'/api/reviews/{rid}/report-fields',json={'expected_version':version,'fields':FIELDS},headers={'X-CSRF-Token':s['csrf_token']}).status_code,404)
        self.assertEqual(self.client.get('/api/session').status_code,200)

    def test_accepted_issue_and_required_final_fields(self):
        r=self.review(self.upload('产品保证收益，风险与费用已披露')['document_id']);rid=r['review_id']
        for i in r['issues']:
            response=self.client.patch(f'/api/reviews/{rid}/issues/{i["issue_id"]}',json=dict(expected_version=r['version'],review_status='accepted',decision_reason='人工核实采纳'),headers=self.headers);r['version']=response.json()['version']
        self.assertEqual(self.client.get(f'/api/reviews/{rid}/report').status_code,409)
        response=self.client.patch(f'/api/reviews/{rid}/report-fields',json={'expected_version':r['version'],'fields':FIELDS},headers=self.headers);r['version']=response.json()['version']
        self.client.post(f'/api/reviews/{rid}/finalize',json=dict(expected_version=r['version'],coverage_confirmed=True,limitations_acknowledged=True),headers=self.headers)
        text=all_text(Document(io.BytesIO(self.client.get(f'/api/reviews/{rid}/report').content)))
        self.assertIn('保证收益',text)

    def test_accepted_ai_candidate_is_in_report_mock_transport_only(self):
        from unittest.mock import AsyncMock, patch
        from pydantic import SecretStr
        from backend.models import AIReviewResult
        text='风险与费用已说明。需要人工判断的表述。'
        d=self.upload(text);anchor='需要人工判断的表述';start=text.index(anchor)
        self.settings.llm_enabled=True;self.settings.api_key=SecretStr('test-only-not-real')
        response=AIReviewResult.model_validate({'issues':[dict(issue_id='model-candidate',risk_type='other',severity='medium',original_text=anchor,location=dict(segment_id='s1',start=start,end=start+len(anchor)),reason='模拟AI候选原因',regulation_ids=['REG-NET-8'],suggestion='模拟AI修改建议',review_status='pending')]})
        with patch('backend.main.AIService.review',new=AsyncMock(return_value=response)):
            r=self.review(d['document_id'],use_ai=True,external_processing_authorized=True)
        self.assertEqual(r['issues'][0]['source'],'ai');self.assertEqual(r['original_issues'][0]['reason'],'模拟AI候选原因')
        rid=r['review_id'];i=r['issues'][0]
        saved=self.client.patch(f'/api/reviews/{rid}/issues/{i["issue_id"]}',json=dict(expected_version=r['version'],review_status='accepted',decision_reason='人工核实采纳'),headers=self.headers)
        meta=self.client.patch(f'/api/reviews/{rid}/report-fields',json={'expected_version':saved.json()['version'],'fields':FIELDS},headers=self.headers)
        final=self.client.post(f'/api/reviews/{rid}/finalize',json=dict(expected_version=meta.json()['version'],coverage_confirmed=True,limitations_acknowledged=True),headers=self.headers)
        self.assertEqual(final.status_code,200,final.text)
        report=all_text(Document(io.BytesIO(self.client.get(f'/api/reviews/{rid}/report').content)))
        self.assertIn('模拟AI候选原因',report);self.assertIn('模拟AI修改建议',report)

    def test_sqlite_saved_review_survives_local_restart_within_ttl(self):
        from backend.main import create_app
        r=self.review(self.upload()['document_id']);rid=r['review_id']
        self.client.patch(f'/api/reviews/{rid}/report-fields',json={'expected_version':r['version'],'fields':FIELDS},headers=self.headers)
        cookies=dict(self.client.cookies);self.client.__exit__(None,None,None)
        self.app=create_app(self.settings);self.client=TestClient(self.app,base_url='https://testserver');self.client.cookies.update(cookies);self.client.__enter__()
        result=self.client.get(f'/api/reviews/{rid}')
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()['report_fields']['reviewer'],FIELDS['reviewer'])
        self.assertEqual(result.json()['original_issues'],r['original_issues'])

    def test_failed_task_retry_creates_new_owned_task(self):
        r=self.review(self.upload()['document_id']);rid=r['review_id']
        # Simulate an interrupted processor record without issuing a real AI request.
        owner=self.app.state.store.db.execute('SELECT owner FROM resources WHERE id=?',(rid,)).fetchone()[0]
        stored=self.app.state.store.get(rid,'review',owner);stored['status']='failed';self.app.state.store.save(rid,stored)
        response=self.client.post(f'/api/reviews/{rid}/retry',json={'expected_version':stored['version']},headers={**self.headers,'Idempotency-Key':'retry-key'})
        self.assertEqual(response.status_code,202,response.text)
        new_id=response.json()['review_id'];self.assertNotEqual(new_id,rid)
        self.assertEqual(self.client.get(f'/api/reviews/{rid}').json()['status'],'failed')
        self.assertEqual(self.client.get(f'/api/reviews/{new_id}').json()['retry_of'],rid)
